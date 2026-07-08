import calendar
import re
import shutil
import sqlite3
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

import openpyxl
from flask import Flask, flash, g, redirect, render_template, request, session, url_for

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
SCHEMA_PATH = BASE_DIR / "schema.sql"
EXPIRING_SOON_DAYS = 60
GRANT_CATEGORIES = {"sponsored", "gift", "internal"}

app = Flask(__name__)
# Only used to sign session/flash cookies for this local, single-user tool (no
# accounts, no sensitive session data) -- a fixed key is fine here.
app.secret_key = "grant-tracker-local-only"
app.jinja_env.filters["money"] = lambda cents: f"{cents / 100:.2f}"
app.jinja_env.filters["rate"] = lambda bps: f"{bps / 100:.2f}"

# Each faculty member gets their own SQLite file under data/, e.g.
# data/jane_smith.db. The active one for this browser session is kept in the
# Flask session so switching faculty doesn't require restarting the app.
NO_FACULTY_REQUIRED_ENDPOINTS = {"faculty_page", "select_faculty", "add_faculty", "static"}


def faculty_db_path(slug):
    return DATA_DIR / f"{slug}.db"


def faculty_backup_path(slug, for_date):
    return DATA_DIR / f"{slug}.db.bak-{for_date:%Y%m%d}"


def list_faculty():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    return sorted(p.stem for p in DATA_DIR.glob("*.db"))


def faculty_display_name(slug):
    return slug.replace("_", " ").title()


def slugify(name):
    return re.sub(r"[^a-z0-9]+", "_", name.strip().lower()).strip("_")


def init_db(path):
    if not path.exists():
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(path)
        db.executescript(SCHEMA_PATH.read_text())
        db.commit()
        db.close()


def migrate_db(db):
    """Add columns introduced after a database file was first created.

    schema.sql only runs for brand-new files (see init_db), so existing
    per-faculty .db files need these added on the fly. Each ALTER is
    idempotent (checked against the live column list) so this is safe to run
    on every connection.
    """
    department_cols = {row["name"] for row in db.execute("PRAGMA table_info(departments)")}
    if "stipend_cents_per_month" not in department_cols:
        db.execute("ALTER TABLE departments ADD COLUMN stipend_cents_per_month INTEGER NOT NULL DEFAULT 0")
    if "tuition_charged_in_summer" not in department_cols:
        db.execute("ALTER TABLE departments ADD COLUMN tuition_charged_in_summer INTEGER NOT NULL DEFAULT 1")
    student_cols = {row["name"] for row in db.execute("PRAGMA table_info(students)")}
    if "expected_graduation" not in student_cols:
        db.execute("ALTER TABLE students ADD COLUMN expected_graduation TEXT")
    if "start_date" not in student_cols:
        db.execute("ALTER TABLE students ADD COLUMN start_date TEXT")
    if "role" not in student_cols:
        db.execute("ALTER TABLE students ADD COLUMN role TEXT NOT NULL DEFAULT 'student'")
    grant_cols = {row["name"] for row in db.execute("PRAGMA table_info(grants)")}
    if "category" not in grant_cols:
        db.execute("ALTER TABLE grants ADD COLUMN category TEXT NOT NULL DEFAULT 'sponsored'")
    has_settings = db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='settings'").fetchone()
    if not has_settings:
        db.execute(
            """CREATE TABLE settings (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                overspend_ratio_bps INTEGER NOT NULL DEFAULT 11500,
                underspend_ratio_bps INTEGER NOT NULL DEFAULT 6000
            )"""
        )
        db.execute("INSERT INTO settings (id, overspend_ratio_bps, underspend_ratio_bps) VALUES (1, 11500, 6000)")
    db.commit()


@app.before_request
def require_faculty():
    if request.endpoint in NO_FACULTY_REQUIRED_ENDPOINTS:
        return None
    slug = session.get("faculty_db")
    if not slug or not faculty_db_path(slug).exists():
        return redirect(url_for("faculty_page"))
    return None


@app.before_request
def backup_before_write():
    """Snapshot the active faculty db before its first write of the day.

    Cheap insurance against a bad edit clobbering real data with no way back (this
    isn't hypothetical -- an errant test write once did exactly that during
    development). One rolling backup per faculty per calendar day, not a full
    history; back the file up yourself before anything riskier (bulk edits, manual
    schema surgery) if you want more than a same-day rollback. Old dated backups
    aren't auto-pruned -- they're plain small SQLite files, so clean them up by hand
    if that ever matters.
    """
    if request.method != "POST":
        return None
    slug = session.get("faculty_db")
    if not slug:
        return None
    db_path = faculty_db_path(slug)
    if not db_path.exists():
        return None
    backup_path = faculty_backup_path(slug, date.today())
    if not backup_path.exists():
        shutil.copy2(db_path, backup_path)
    return None


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(faculty_db_path(session["faculty_db"]))
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
        migrate_db(g.db)
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def parse_money(raw):
    """Parse a user-entered dollar amount into non-negative integer cents, or None if invalid."""
    try:
        value = Decimal(raw.strip())
    except (InvalidOperation, AttributeError):
        return None
    if value < 0:
        return None
    return int((value * 100).quantize(Decimal("1")))


def parse_rate(raw):
    """Parse a user-entered percentage (e.g. '54.5') into non-negative integer basis points, or None if invalid."""
    try:
        value = Decimal(raw.strip())
    except (InvalidOperation, AttributeError):
        return None
    if value < 0:
        return None
    return int((value * 100).quantize(Decimal("1")))


def parse_date(raw):
    try:
        return date.fromisoformat(raw.strip())
    except (ValueError, AttributeError):
        return None


def parse_category(raw):
    raw = (raw or "").strip().lower()
    return raw if raw in GRANT_CATEGORIES else None


def parse_optional_date(raw):
    """Like parse_date, but a blank/missing value is valid and means 'not set'.

    Returns (date_or_None, ok) -- ok is False only when raw is non-blank but not a valid date.
    """
    raw = (raw or "").strip()
    if not raw:
        return None, True
    parsed = parse_date(raw)
    return parsed, parsed is not None


def parse_month(raw):
    """Parse a 'YYYY-MM' string (from an <input type=month>) into the first-of-month date."""
    try:
        year_str, month_str = raw.strip().split("-")
        return date(int(year_str), int(month_str), 1)
    except (ValueError, AttributeError):
        return None


def month_range(start, end):
    months = []
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        months.append(f"{year:04d}-{month:02d}")
        month += 1
        if month > 12:
            month = 1
            year += 1
    return months


def month_label(month_str):
    year, month = month_str.split("-")
    return f"{calendar.month_abbr[int(month)]} {year}"


def month_end(month_start):
    last_day = calendar.monthrange(month_start.year, month_start.month)[1]
    return date(month_start.year, month_start.month, last_day)


def is_summer_month(month_str):
    return int(month_str.split("-")[1]) in (6, 7, 8)


def parse_scenario_arg(raw):
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def scenario_clause(scenario_id):
    if scenario_id is None:
        return "scenario_id IS NULL", ()
    return "scenario_id = ?", (scenario_id,)


def grant_status(end_date_str):
    days_left = (date.fromisoformat(end_date_str) - date.today()).days
    if days_left < 0:
        return "expired"
    if days_left <= EXPIRING_SOON_DAYS:
        return "expiring-soon"
    return "active"


def is_chargeable(month_str, start_date_str, expected_graduation_str):
    """False if `month_str` (YYYY-MM) falls entirely outside [start_date, expected_graduation].

    A student isn't charged to any grant for months entirely before their start date or
    entirely after their expected graduation. A month that only partially overlaps that
    window (they start or graduate mid-month) is still charged in full, since allocations
    are tracked at monthly granularity.
    """
    month_start = parse_month(month_str)
    if expected_graduation_str and month_start > date.fromisoformat(expected_graduation_str):
        return False
    if start_date_str and month_end(month_start) < date.fromisoformat(start_date_str):
        return False
    return True


def allocation_cost_cents(stipend_cents_per_month, tuition_cents_per_month, fringe_rate_bps, overhead_rate_bps, percent):
    """Monthly cost breakdown for allocating a student at `percent`% to a grant.

    Tuition is prorated by effort % but excluded from the overhead base, matching how
    federal MTDC (modified total direct cost) typically excludes tuition remission.
    """
    stipend = stipend_cents_per_month * percent // 100
    tuition = tuition_cents_per_month * percent // 100
    fringe = stipend * fringe_rate_bps // 10000
    overhead = (stipend + fringe) * overhead_rate_bps // 10000
    return {
        "stipend": stipend,
        "tuition": tuition,
        "fringe": fringe,
        "overhead": overhead,
        "total": stipend + tuition + fringe + overhead,
    }


def empty_cost_breakdown():
    return {"stipend": 0, "tuition": 0, "fringe": 0, "overhead": 0, "total": 0}


def add_cost_breakdown(totals, addition):
    for key in totals:
        totals[key] += addition[key]


def grant_with_balance(row):
    spent_cents = get_db().execute(
        "SELECT COALESCE(SUM(amount_cents), 0) AS total FROM transactions WHERE grant_id = ?",
        (row["id"],),
    ).fetchone()["total"]
    grant = dict(row)
    grant["spent_cents"] = spent_cents
    grant["balance_cents"] = row["total_amount_cents"] - spent_cents
    grant["status"] = grant_status(row["end_date"])
    return grant


def current_monthly_burn_cents(grid):
    """Projected personnel cost for the grant's "current" month: this calendar month if
    it's in the allocation grid, else the nearest future month, else 0 if there's no
    ongoing/future allocation to extrapolate a burn rate from."""
    if not grid:
        return 0
    today_month = date.today().strftime("%Y-%m")
    for month, cost in zip(grid["months"], grid["monthly_costs"]):
        if month >= today_month:
            return cost
    return 0


def get_settings(db):
    return dict(db.execute("SELECT overspend_ratio_bps, underspend_ratio_bps FROM settings WHERE id = 1").fetchone())


def grant_spending_risk(grant, grid, overspend_ratio_bps, underspend_ratio_bps):
    """'overspending', 'underspending', or None (on track / not evaluated).

    Extrapolates the grant's current monthly personnel burn rate (see
    current_monthly_burn_cents) across its remaining months and compares that
    projection to its remaining (unspent) balance. Only evaluated for non-expired
    grants -- an expired grant's spending is done, not "at risk". The thresholds
    are user-tunable (see the Settings page / get_settings) rather than fixed
    constants, since how sensitive this should be varies by grant portfolio.
    """
    if grant["status"] == "expired":
        return None
    balance = grant["balance_cents"]
    if balance <= 0:
        return "overspending"

    burn_rate = current_monthly_burn_cents(grid)
    if burn_rate == 0:
        return "underspending"

    today = date.today()
    end_date = date.fromisoformat(grant["end_date"])
    months_remaining = len(month_range(date(today.year, today.month, 1), date(end_date.year, end_date.month, 1)))
    projected_remaining_spend = burn_rate * months_remaining

    ratio = projected_remaining_spend / balance
    if ratio > overspend_ratio_bps / 10000:
        return "overspending"
    if ratio < underspend_ratio_bps / 10000:
        return "underspending"
    return None


def grant_allocation_grid(grant_id, scenario_id):
    """Students x months grid of % allocation to this grant, plus projected cost breakdown."""
    db = get_db()
    grant_row = db.execute("SELECT overhead_rate_bps FROM grants WHERE id = ?", (grant_id,)).fetchone()
    if grant_row is None:
        return None
    overhead_rate_bps = grant_row["overhead_rate_bps"]

    clause, params = scenario_clause(scenario_id)
    rows = db.execute(
        f"""SELECT allocations.month, allocations.percent, students.id AS student_id,
                   students.name AS student_name, students.stipend_cents_per_month,
                   students.start_date, students.expected_graduation,
                   COALESCE(departments.tuition_cents_per_month, 0) AS tuition_cents_per_month,
                   COALESCE(departments.fringe_rate_bps, 0) AS fringe_rate_bps,
                   COALESCE(departments.tuition_charged_in_summer, 1) AS tuition_charged_in_summer
            FROM allocations
            JOIN students ON students.id = allocations.student_id
            LEFT JOIN departments ON departments.id = students.department_id
            WHERE allocations.grant_id = ? AND {clause}""",
        (grant_id, *params),
    ).fetchall()
    if not rows:
        return None

    months = month_range(
        min(parse_month(r["month"]) for r in rows), max(parse_month(r["month"]) for r in rows)
    )

    by_student = {}
    for r in rows:
        entry = by_student.setdefault(
            r["student_id"],
            {
                "name": r["student_name"],
                "start_date": r["start_date"],
                "expected_graduation": r["expected_graduation"],
                "percents": {},
            },
        )
        entry["percents"][r["month"]] = r["percent"]

    student_rows = [
        {
            "name": data["name"],
            "cells": [
                {
                    "percent": data["percents"].get(m),
                    "chargeable": is_chargeable(m, data["start_date"], data["expected_graduation"]),
                }
                for m in months
            ],
        }
        for _, data in sorted(by_student.items(), key=lambda kv: kv[1]["name"])
    ]

    cost_by_month = {m: empty_cost_breakdown() for m in months}
    total_breakdown = empty_cost_breakdown()
    for r in rows:
        if is_chargeable(r["month"], r["start_date"], r["expected_graduation"]):
            tuition_cents = r["tuition_cents_per_month"]
            if not r["tuition_charged_in_summer"] and is_summer_month(r["month"]):
                tuition_cents = 0
            cost = allocation_cost_cents(
                r["stipend_cents_per_month"], tuition_cents, r["fringe_rate_bps"], overhead_rate_bps, r["percent"]
            )
        else:
            cost = empty_cost_breakdown()
        add_cost_breakdown(cost_by_month[r["month"]], cost)
        add_cost_breakdown(total_breakdown, cost)
    monthly_costs = [cost_by_month[m]["total"] for m in months]

    return {
        "months": months,
        "month_labels": [month_label(m) for m in months],
        "rows": student_rows,
        "monthly_costs": monthly_costs,
        "total_cost": sum(monthly_costs),
        "breakdown": total_breakdown,
    }


def student_allocation_grid(student_id, scenario_id):
    """Grants x months grid of % allocation for this student, plus a per-month total (should stay <=100)."""
    db = get_db()
    student_row = db.execute(
        """SELECT students.stipend_cents_per_month, students.start_date, students.expected_graduation,
                  COALESCE(departments.tuition_cents_per_month, 0) AS tuition_cents_per_month,
                  COALESCE(departments.fringe_rate_bps, 0) AS fringe_rate_bps,
                  COALESCE(departments.tuition_charged_in_summer, 1) AS tuition_charged_in_summer
           FROM students LEFT JOIN departments ON departments.id = students.department_id
           WHERE students.id = ?""",
        (student_id,),
    ).fetchone()
    if student_row is None:
        return None

    clause, params = scenario_clause(scenario_id)
    rows = db.execute(
        f"""SELECT allocations.month, allocations.percent, grants.id AS grant_id, grants.name AS grant_name,
                   grants.overhead_rate_bps
            FROM allocations JOIN grants ON grants.id = allocations.grant_id
            WHERE allocations.student_id = ? AND {clause}""",
        (student_id, *params),
    ).fetchall()
    if not rows:
        return None

    months = month_range(
        min(parse_month(r["month"]) for r in rows), max(parse_month(r["month"]) for r in rows)
    )

    by_grant = {}
    for r in rows:
        entry = by_grant.setdefault(r["grant_id"], {"name": r["grant_name"], "percents": {}})
        entry["percents"][r["month"]] = r["percent"]

    grant_rows = [
        {"grant_id": grant_id, "name": data["name"], "cells": [data["percents"].get(m) for m in months]}
        for grant_id, data in sorted(by_grant.items(), key=lambda kv: kv[1]["name"])
    ]

    totals = [0] * len(months)
    for row in grant_rows:
        for i, cell in enumerate(row["cells"]):
            if cell:
                totals[i] += cell

    months_chargeable = [is_chargeable(m, student_row["start_date"], student_row["expected_graduation"]) for m in months]

    total_breakdown = empty_cost_breakdown()
    for r in rows:
        if is_chargeable(r["month"], student_row["start_date"], student_row["expected_graduation"]):
            tuition_cents = student_row["tuition_cents_per_month"]
            if not student_row["tuition_charged_in_summer"] and is_summer_month(r["month"]):
                tuition_cents = 0
            cost = allocation_cost_cents(
                student_row["stipend_cents_per_month"],
                tuition_cents,
                student_row["fringe_rate_bps"],
                r["overhead_rate_bps"],
                r["percent"],
            )
        else:
            cost = empty_cost_breakdown()
        add_cost_breakdown(total_breakdown, cost)

    return {
        "months": months,
        "month_labels": [month_label(m) for m in months],
        "months_chargeable": months_chargeable,
        "rows": grant_rows,
        "totals": totals,
        "breakdown": total_breakdown,
    }


def under_allocated_months(grid):
    """[(month_label, total_percent), ...] for chargeable months in a student_allocation_grid
    result where the allocated total is under 100%. [] if grid is None (nothing allocated) or
    nothing's under. Used to flag students who aren't fully committed for some month."""
    if not grid:
        return []
    return [
        (grid["month_labels"][i], total)
        for i, total in enumerate(grid["totals"])
        if grid["months_chargeable"][i] and total < 100
    ]


# --- "Plan a new hire" capacity report (see plan_new_student()) ---------------------


def grant_spare_capacity_cents(grant, months):
    """Grant's remaining balance minus its already-committed (live) projected personnel
    cost for `months`. Does NOT reserve any balance for commitments outside `months` --
    a grant heavily committed just past the window's edge can still show its full
    uncommitted balance as spare here."""
    grid = grant_allocation_grid(grant["id"], None)
    committed_by_month = dict(zip(grid["months"], grid["monthly_costs"])) if grid else {}
    committed = sum(committed_by_month.get(m, 0) for m in months)
    return grant["balance_cents"] - committed


def eligible_grants_for_window(db, months):
    """(eligible, excluded) grants for a capacity-planning window. `eligible` is
    non-expired grants whose [start_date, end_date] fully covers the window and that
    still have positive spare_cents left in it (each gets a "spare_cents" key added).
    `excluded` is [(grant, reason), ...] for everything else, for a transparency panel."""
    window_start = parse_month(months[0])
    window_end_date = month_end(parse_month(months[-1]))
    eligible, excluded = [], []
    for row in db.execute("SELECT * FROM grants ORDER BY name"):
        grant = grant_with_balance(row)
        if grant["status"] == "expired":
            excluded.append((grant, "expired"))
            continue
        if date.fromisoformat(grant["start_date"]) > window_start or date.fromisoformat(grant["end_date"]) < window_end_date:
            excluded.append((grant, "doesn't cover the full window"))
            continue
        spare_cents = grant_spare_capacity_cents(grant, months)
        if spare_cents <= 0:
            excluded.append((grant, "no spare capacity left in this window"))
            continue
        grant["spare_cents"] = spare_cents
        eligible.append(grant)
    return eligible, excluded


def hire_window_cost_cents(department, overhead_rate_bps, months):
    """([cost_cents_per_month, ...], total) for one hypothetical 100%-effort hire from
    `department`, funded by a grant with `overhead_rate_bps`, across `months`. Applies
    the summer-tuition-exclusion rule the same way the live grids do."""
    monthly = []
    for m in months:
        tuition_cents = department["tuition_cents_per_month"]
        if not department["tuition_charged_in_summer"] and is_summer_month(m):
            tuition_cents = 0
        cost = allocation_cost_cents(
            department["stipend_cents_per_month"], tuition_cents, department["fringe_rate_bps"], overhead_rate_bps, 100
        )
        monthly.append(cost["total"])
    return monthly, sum(monthly)


def available_percent(spare_cents_remaining, window_cost_100pct_cents):
    """How much % of one hypothetical 100%-effort hire's window cost the remaining spare
    dollars could still fund, clamped to [0, 100]. A zero-cost hire (e.g. all-zero
    department rates) is trivially 100% fundable rather than a division by zero."""
    if window_cost_100pct_cents <= 0:
        return 100
    return max(0, min(100, int(spare_cents_remaining * 100 // window_cost_100pct_cents)))


def _new_hire_result(department, role, assignments, placed_pct, months):
    window_cost_cents = sum(a["window_cost_cents"] for a in assignments)
    return {
        "department_name": department["name"],
        "role": role,
        "assignments": assignments,
        "placed_pct": placed_pct,
        "feasible": placed_pct >= 100,
        "window_cost_cents": window_cost_cents,
        "monthly_cost_cents": window_cost_cents / len(months) if months else 0,
    }


def _new_hire_plan_result(name, description, hire_results):
    fully_placed = sum(1 for h in hire_results if h["feasible"])
    return {
        "name": name,
        "description": description,
        "hires": hire_results,
        "fully_placed_count": fully_placed,
        "requested_count": len(hire_results),
        "feasible": fully_placed == len(hire_results),
        "total_cost_cents": sum(h["window_cost_cents"] for h in hire_results),
    }


def build_greedy_hire_plan(name, description, hires, grant_order, months):
    """Fills each hire's 100% need by walking `grant_order` (already sorted by the
    caller -- most-spare-first for "Concentrate", soonest-expiring-first for "Use it or
    lose it") and draining a shared spare_remaining pool per grant. Because the pool
    persists across hires against one fixed order, a grant is naturally exhausted before
    later hires spill onto the next grant in the order -- no special-casing needed."""
    spare_remaining = {g["id"]: float(g["spare_cents"]) for g in grant_order}
    cost_cache = {}
    hire_results = []
    for hire in hires:
        dept = hire["department"]
        remaining_need = 100
        assignments = []
        for grant in grant_order:
            if remaining_need <= 0:
                break
            key = (grant["id"], dept["id"])
            if key not in cost_cache:
                cost_cache[key] = hire_window_cost_cents(dept, grant["overhead_rate_bps"], months)
            _, cost100 = cost_cache[key]
            take_pct = min(remaining_need, available_percent(spare_remaining[grant["id"]], cost100))
            if take_pct <= 0:
                continue
            cost_for_take = cost100 * take_pct / 100
            spare_remaining[grant["id"]] -= cost_for_take
            remaining_need -= take_pct
            assignments.append(
                {"grant_id": grant["id"], "grant_name": grant["name"], "percent": take_pct, "window_cost_cents": cost_for_take}
            )
        hire_results.append(_new_hire_result(dept, hire["role"], assignments, 100 - remaining_need, months))
    return _new_hire_plan_result(name, description, hire_results)


def build_spread_hire_plan(hires, grants, months):
    """For each hire (in submitted order), splits its 100% need proportionally across
    every grant with remaining spare capacity, using largest-remainder rounding so
    per-grant percentages are integers that sum exactly to the feasible total. Recomputed
    fresh per hire against the shared remaining pool, so later hires reflect what earlier
    ones already consumed."""
    spare_remaining = {g["id"]: float(g["spare_cents"]) for g in grants}
    cost_cache = {}
    hire_results = []
    for hire in hires:
        dept = hire["department"]
        candidates = []
        for grant in grants:
            if spare_remaining[grant["id"]] <= 0:
                continue
            key = (grant["id"], dept["id"])
            if key not in cost_cache:
                cost_cache[key] = hire_window_cost_cents(dept, grant["overhead_rate_bps"], months)
            _, cost100 = cost_cache[key]
            avail_pct = available_percent(spare_remaining[grant["id"]], cost100)
            if avail_pct > 0:
                candidates.append((grant, avail_pct, cost100))

        total_avail = sum(avail_pct for _, avail_pct, _ in candidates)
        if total_avail == 0:
            hire_results.append(_new_hire_result(dept, hire["role"], [], 0, months))
            continue

        target_total = min(100, total_avail)
        raw = [(grant, cost100, target_total * avail_pct / total_avail) for grant, avail_pct, cost100 in candidates]
        floors = [(grant, cost100, int(t), t - int(t)) for grant, cost100, t in raw]
        assigned_sum = sum(f for _, _, f, _ in floors)
        remainder_needed = target_total - assigned_sum
        floors.sort(key=lambda row: -row[3])

        assignments = []
        for i, (grant, cost100, f, _frac) in enumerate(floors):
            pct = f + (1 if i < remainder_needed else 0)
            if pct <= 0:
                continue
            cost_for_take = cost100 * pct / 100
            spare_remaining[grant["id"]] -= cost_for_take
            assignments.append(
                {"grant_id": grant["id"], "grant_name": grant["name"], "percent": pct, "window_cost_cents": cost_for_take}
            )
        placed_pct = sum(a["percent"] for a in assignments)
        hire_results.append(_new_hire_result(dept, hire["role"], assignments, placed_pct, months))

    return _new_hire_plan_result(
        "Spread evenly",
        "Splits each hire's effort proportionally across every grant with remaining spare capacity.",
        hire_results,
    )


@app.route("/")
def index():
    db = get_db()
    category_filter = set(request.args.getlist("category")) & GRANT_CATEGORIES
    hide_closed = request.args.get("hide_closed") == "1"

    query = "SELECT * FROM grants"
    params = ()
    if category_filter:
        placeholders = ",".join("?" for _ in category_filter)
        query += f" WHERE category IN ({placeholders})"
        params = tuple(category_filter)
    query += " ORDER BY end_date"

    grants = [grant_with_balance(r) for r in db.execute(query, params)]
    if hide_closed:
        grants = [g for g in grants if g["status"] != "expired" and g["balance_cents"] >= 0]
    settings = get_settings(db)
    for grant in grants:
        grid = grant_allocation_grid(grant["id"], None)
        grant["projected_personnel_cents"] = grid["total_cost"] if grid else 0
        grant["spending_risk"] = grant_spending_risk(
            grant, grid, settings["overspend_ratio_bps"], settings["underspend_ratio_bps"]
        )
    students = db.execute(
        """SELECT students.*, departments.name AS department_name FROM students
           LEFT JOIN departments ON departments.id = students.department_id
           ORDER BY students.name"""
    ).fetchall()
    departments = db.execute("SELECT * FROM departments ORDER BY name").fetchall()
    # Independent of the category/hide-closed filters above -- the "initial allocation"
    # editor on the add-student form should offer every grant regardless of what the
    # dashboard's grant table currently has filtered out.
    all_grants_json = [
        {"id": r["id"], "name": r["name"]} for r in db.execute("SELECT id, name FROM grants ORDER BY name")
    ]

    # Each category pill's link toggles just that category in/out of the current
    # selection (keeping the others and hide_closed), so more than one can be active
    # at once; the "All" pill clears the selection entirely.
    hide_closed_param = 1 if hide_closed else None
    category_toggle_urls = {
        category: url_for("index", category=sorted(category_filter ^ {category}), hide_closed=hide_closed_param)
        for category in GRANT_CATEGORIES
    }
    category_clear_url = url_for("index", hide_closed=hide_closed_param)

    hide_closed_toggle_url = url_for("index", category=sorted(category_filter), hide_closed=None if hide_closed else 1)

    return render_template(
        "index.html",
        grants=grants,
        students=students,
        departments=departments,
        today=date.today().isoformat(),
        category_filter=category_filter,
        category_toggle_urls=category_toggle_urls,
        category_clear_url=category_clear_url,
        hide_closed=hide_closed,
        hide_closed_toggle_url=hide_closed_toggle_url,
        grants_json=all_grants_json,
        baseline_allocations_by_student={},
    )


@app.route("/grants/add", methods=["POST"])
def add_grant():
    name = request.form.get("name", "").strip()
    total_amount_cents = parse_money(request.form.get("total_amount", ""))
    start_date = parse_date(request.form.get("start_date", ""))
    end_date = parse_date(request.form.get("end_date", ""))
    overhead_rate_bps = parse_rate(request.form.get("overhead_rate", "0") or "0")
    category = parse_category(request.form.get("category", "sponsored"))

    if not name:
        flash("Grant name is required.")
    elif total_amount_cents is None:
        flash("Total amount must be a non-negative number.")
    elif start_date is None or end_date is None:
        flash("Start and end dates must be valid dates.")
    elif end_date < start_date:
        flash("End date can't be before the start date.")
    elif overhead_rate_bps is None:
        flash("Overhead rate must be a non-negative number.")
    elif category is None:
        flash("Category must be sponsored, gift, or internal.")
    else:
        get_db().execute(
            """INSERT INTO grants (name, sponsor, total_amount_cents, start_date, end_date, overhead_rate_bps,
               category, notes) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                name,
                request.form.get("sponsor", "").strip(),
                total_amount_cents,
                start_date.isoformat(),
                end_date.isoformat(),
                overhead_rate_bps,
                category,
                request.form.get("notes", "").strip(),
            ),
        )
        get_db().commit()
    return redirect(url_for("index"))


@app.route("/grants/<int:grant_id>")
def grant_detail(grant_id):
    db = get_db()
    row = db.execute("SELECT * FROM grants WHERE id = ?", (grant_id,)).fetchone()
    if row is None:
        return redirect(url_for("index"))

    scenario_id = parse_scenario_arg(request.args.get("scenario"))
    scenarios = db.execute("SELECT * FROM scenarios ORDER BY created_at DESC").fetchall()
    students = db.execute("SELECT * FROM students ORDER BY name").fetchall()
    transactions = db.execute(
        "SELECT * FROM transactions WHERE grant_id = ? ORDER BY date DESC, id DESC", (grant_id,)
    ).fetchall()

    grid = grant_allocation_grid(grant_id, scenario_id)
    live_grid = grant_allocation_grid(grant_id, None)
    live_cost_cents = live_grid["total_cost"] if live_grid else 0
    scenario_cost_cents = grid["total_cost"] if (grid and scenario_id is not None) else None

    settings = get_settings(db)
    grant = grant_with_balance(row)
    grant["spending_risk"] = grant_spending_risk(
        grant, live_grid, settings["overspend_ratio_bps"], settings["underspend_ratio_bps"]
    )

    return render_template(
        "grant_detail.html",
        grant=grant,
        transactions=transactions,
        grid=grid,
        students=students,
        scenarios=scenarios,
        current_scenario_id=scenario_id,
        live_cost_cents=live_cost_cents,
        scenario_cost_cents=scenario_cost_cents,
        today=date.today().isoformat(),
    )


@app.route("/grants/<int:grant_id>/edit", methods=["POST"])
def edit_grant(grant_id):
    name = request.form.get("name", "").strip()
    total_amount_cents = parse_money(request.form.get("total_amount", ""))
    start_date = parse_date(request.form.get("start_date", ""))
    end_date = parse_date(request.form.get("end_date", ""))
    overhead_rate_bps = parse_rate(request.form.get("overhead_rate", "0") or "0")
    category = parse_category(request.form.get("category", "sponsored"))

    if not name:
        flash("Grant name is required.")
    elif total_amount_cents is None:
        flash("Total amount must be a non-negative number.")
    elif start_date is None or end_date is None:
        flash("Start and end dates must be valid dates.")
    elif end_date < start_date:
        flash("End date can't be before the start date.")
    elif overhead_rate_bps is None:
        flash("Overhead rate must be a non-negative number.")
    elif category is None:
        flash("Category must be sponsored, gift, or internal.")
    else:
        get_db().execute(
            """UPDATE grants SET name = ?, sponsor = ?, total_amount_cents = ?, start_date = ?, end_date = ?,
               overhead_rate_bps = ?, category = ?, notes = ? WHERE id = ?""",
            (
                name,
                request.form.get("sponsor", "").strip(),
                total_amount_cents,
                start_date.isoformat(),
                end_date.isoformat(),
                overhead_rate_bps,
                category,
                request.form.get("notes", "").strip(),
                grant_id,
            ),
        )
        get_db().commit()
    return redirect(url_for("grant_detail", grant_id=grant_id))


@app.route("/grants/<int:grant_id>/delete", methods=["POST"])
def delete_grant(grant_id):
    db = get_db()
    in_use = db.execute(
        """SELECT 1 FROM transactions WHERE grant_id = ?
           UNION SELECT 1 FROM allocations WHERE grant_id = ? LIMIT 1""",
        (grant_id, grant_id),
    ).fetchone()
    if in_use:
        flash("Can't delete a grant that has transactions or allocations (in live data or a scenario). Remove those first.")
        return redirect(url_for("grant_detail", grant_id=grant_id))
    db.execute("DELETE FROM grants WHERE id = ?", (grant_id,))
    db.commit()
    return redirect(url_for("index"))


@app.route("/grants/<int:grant_id>/transactions/add", methods=["POST"])
def add_transaction(grant_id):
    tx_date = parse_date(request.form.get("date", ""))
    amount_cents = parse_money(request.form.get("amount", ""))

    if tx_date is None:
        flash("Transaction date must be valid.")
    elif amount_cents is None:
        flash("Transaction amount must be a non-negative number.")
    else:
        get_db().execute(
            "INSERT INTO transactions (grant_id, date, amount_cents, description) VALUES (?, ?, ?, ?)",
            (grant_id, tx_date.isoformat(), amount_cents, request.form.get("description", "").strip()),
        )
        get_db().commit()
    return redirect(url_for("grant_detail", grant_id=grant_id))


@app.route("/transactions/<int:transaction_id>/edit", methods=["POST"])
def edit_transaction(transaction_id):
    db = get_db()
    row = db.execute("SELECT grant_id FROM transactions WHERE id = ?", (transaction_id,)).fetchone()
    if row is None:
        return redirect(url_for("index"))
    grant_id = row["grant_id"]

    tx_date = parse_date(request.form.get("date", ""))
    amount_cents = parse_money(request.form.get("amount", ""))

    if tx_date is None:
        flash("Transaction date must be valid.")
    elif amount_cents is None:
        flash("Transaction amount must be a non-negative number.")
    else:
        db.execute(
            "UPDATE transactions SET date = ?, amount_cents = ?, description = ? WHERE id = ?",
            (tx_date.isoformat(), amount_cents, request.form.get("description", "").strip(), transaction_id),
        )
        db.commit()
    return redirect(url_for("grant_detail", grant_id=grant_id))


@app.route("/transactions/<int:transaction_id>/delete", methods=["POST"])
def delete_transaction(transaction_id):
    db = get_db()
    row = db.execute("SELECT grant_id FROM transactions WHERE id = ?", (transaction_id,)).fetchone()
    if row is not None:
        db.execute("DELETE FROM transactions WHERE id = ?", (transaction_id,))
        db.commit()
        return redirect(url_for("grant_detail", grant_id=row["grant_id"]))
    return redirect(url_for("index"))


def parse_department_id(raw):
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


STUDENT_ROLES = {"student", "postdoc"}


def parse_role(raw):
    raw = (raw or "").strip().lower()
    return raw if raw in STUDENT_ROLES else None


@app.route("/students/add", methods=["POST"])
def add_student():
    name = request.form.get("name", "").strip()
    stipend_cents = parse_money(request.form.get("stipend", "0") or "0")
    department_id = parse_department_id(request.form.get("department_id"))
    role = parse_role(request.form.get("role", "student"))
    start_date, start_date_ok = parse_optional_date(request.form.get("start_date", ""))
    expected_graduation, graduation_ok = parse_optional_date(request.form.get("expected_graduation", ""))
    if not name:
        flash("Student name is required.")
    elif stipend_cents is None:
        flash("Monthly stipend must be a non-negative number.")
    elif role is None:
        flash("Role must be student or postdoc.")
    elif not start_date_ok:
        flash("Start date must be a valid date.")
    elif not graduation_ok:
        flash("Expected graduation must be a valid date.")
    else:
        db = get_db()
        cursor = db.execute(
            """INSERT INTO students (name, email, department_id, role, stipend_cents_per_month, start_date,
               expected_graduation, notes) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                name,
                request.form.get("email", "").strip(),
                department_id,
                role,
                stipend_cents,
                start_date.isoformat() if start_date else None,
                expected_graduation.isoformat() if expected_graduation else None,
                request.form.get("notes", "").strip(),
            ),
        )
        new_student_id = cursor.lastrowid

        grant_ids = request.form.getlist("grant_id[]")
        percents = request.form.getlist("percent[]")
        rows = [(g, p) for g, p in zip(grant_ids, percents) if g]
        if rows:
            error = apply_allocation_batch(
                db,
                None,
                new_student_id,
                rows,
                request.form.get("month_start", ""),
                request.form.get("month_end", ""),
            )
            if error:
                db.rollback()
                flash(error)
                return redirect(url_for("index"))

        db.commit()
    return redirect(url_for("index"))


@app.route("/students/<int:student_id>")
def student_detail(student_id):
    db = get_db()
    student = db.execute(
        """SELECT students.*, departments.name AS department_name FROM students
           LEFT JOIN departments ON departments.id = students.department_id
           WHERE students.id = ?""",
        (student_id,),
    ).fetchone()
    if student is None:
        return redirect(url_for("index"))

    scenario_id = parse_scenario_arg(request.args.get("scenario"))
    scenarios = db.execute("SELECT * FROM scenarios ORDER BY created_at DESC").fetchall()
    grants = db.execute("SELECT * FROM grants ORDER BY name").fetchall()
    departments = db.execute("SELECT * FROM departments ORDER BY name").fetchall()
    grid = student_allocation_grid(student_id, scenario_id)

    return render_template(
        "student_detail.html",
        student=student,
        grants=grants,
        departments=departments,
        scenarios=scenarios,
        current_scenario_id=scenario_id,
        grid=grid,
        under_allocated=under_allocated_months(grid),
        today=date.today().isoformat(),
    )


@app.route("/students/<int:student_id>/edit", methods=["POST"])
def edit_student(student_id):
    name = request.form.get("name", "").strip()
    stipend_cents = parse_money(request.form.get("stipend", "0") or "0")
    department_id = parse_department_id(request.form.get("department_id"))
    role = parse_role(request.form.get("role", "student"))
    start_date, start_date_ok = parse_optional_date(request.form.get("start_date", ""))
    expected_graduation, graduation_ok = parse_optional_date(request.form.get("expected_graduation", ""))
    if not name:
        flash("Student name is required.")
    elif stipend_cents is None:
        flash("Monthly stipend must be a non-negative number.")
    elif role is None:
        flash("Role must be student or postdoc.")
    elif not start_date_ok:
        flash("Start date must be a valid date.")
    elif not graduation_ok:
        flash("Expected graduation must be a valid date.")
    else:
        get_db().execute(
            """UPDATE students SET name = ?, email = ?, department_id = ?, role = ?, stipend_cents_per_month = ?,
               start_date = ?, expected_graduation = ?, notes = ? WHERE id = ?""",
            (
                name,
                request.form.get("email", "").strip(),
                department_id,
                role,
                stipend_cents,
                start_date.isoformat() if start_date else None,
                expected_graduation.isoformat() if expected_graduation else None,
                request.form.get("notes", "").strip(),
                student_id,
            ),
        )
        get_db().commit()
    return redirect(url_for("student_detail", student_id=student_id))


@app.route("/students/<int:student_id>/delete", methods=["POST"])
def delete_student(student_id):
    db = get_db()
    in_use = db.execute("SELECT 1 FROM allocations WHERE student_id = ? LIMIT 1", (student_id,)).fetchone()
    if in_use:
        flash("Can't delete a student who has allocations (in live data or a scenario). Remove those first.")
        return redirect(url_for("student_detail", student_id=student_id))
    db.execute("DELETE FROM students WHERE id = ?", (student_id,))
    db.commit()
    return redirect(url_for("index"))


def apply_allocation(db, scenario_id, student_id, grant_id, month_start_raw, month_end_raw, percent_raw):
    """Validate and write a student/grant/month-range/percent allocation change.

    Returns an error message string on failure, or None on success. Used by both the
    allocate forms on grant/student pages and the "set an initial change" step when
    creating a new what-if scenario, so the same over-100% validation applies everywhere.
    """
    student = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    grant_exists = db.execute("SELECT 1 FROM grants WHERE id = ?", (grant_id,)).fetchone()
    month_start = parse_month(month_start_raw)
    month_end = parse_month(month_end_raw)
    try:
        percent = int(percent_raw)
    except (TypeError, ValueError):
        percent = None

    if student is None or not grant_exists:
        return "Select a valid student and grant."
    if month_start is None or month_end is None:
        return "Start and end month must be valid."
    if month_end < month_start:
        return "End month can't be before the start month."
    if percent is None or not (0 <= percent <= 100):
        return "Percent must be a whole number between 0 and 100."

    months = month_range(month_start, month_end)
    clause, params = scenario_clause(scenario_id)

    for month in months:
        other_total = db.execute(
            f"""SELECT COALESCE(SUM(percent), 0) AS total FROM allocations
                WHERE student_id = ? AND month = ? AND grant_id != ? AND {clause}""",
            (student_id, month, grant_id, *params),
        ).fetchone()["total"]
        if other_total + percent > 100:
            return f"{student['name']} would be over 100% allocated in {month_label(month)} ({other_total + percent}%)."

    for month in months:
        db.execute(
            f"DELETE FROM allocations WHERE student_id = ? AND grant_id = ? AND month = ? AND {clause}",
            (student_id, grant_id, month, *params),
        )
        if percent > 0:
            db.execute(
                "INSERT INTO allocations (scenario_id, student_id, grant_id, month, percent) VALUES (?, ?, ?, ?, ?)",
                (scenario_id, student_id, grant_id, month, percent),
            )
    return None


def apply_allocation_batch(db, scenario_id, student_id, grant_percent_pairs, month_start_raw, month_end_raw):
    """Validate and write several (grant_id, percent) allocations for one student over a
    month range, all at once.

    Unlike apply_allocation (one grant at a time, validated against whatever is already
    in the database), this validates the *combined* new total up front before writing
    anything. That matters for rebalancing: reducing grant A from 100% to 50% while
    raising grant B from 0% to 50% would fail with apply_allocation if you raised B
    first (A is still at 100% at that instant) -- here both changes are checked together
    against each other, plus any grant not included in the batch, and only written if
    the whole batch keeps every month at <=100%.

    Returns an error message string on failure, or None on success. Rows with percent 0
    clear that grant's allocation for the range, same as apply_allocation.
    """
    student = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    month_start = parse_month(month_start_raw)
    month_end = parse_month(month_end_raw)
    if student is None:
        return "Select a valid student."
    if month_start is None or month_end is None:
        return "Start and end month must be valid."
    if month_end < month_start:
        return "End month can't be before the start month."

    rows = []
    for grant_id_raw, percent_raw in grant_percent_pairs:
        try:
            grant_id = int(grant_id_raw)
            percent = int(percent_raw)
        except (TypeError, ValueError):
            return "Grant and percent must be valid."
        if not (0 <= percent <= 100):
            return "Percent must be a whole number between 0 and 100."
        if not db.execute("SELECT 1 FROM grants WHERE id = ?", (grant_id,)).fetchone():
            return "Select a valid grant."
        rows.append((grant_id, percent))

    grant_ids = [grant_id for grant_id, _ in rows]
    batch_total = sum(percent for _, percent in rows)
    months = month_range(month_start, month_end)
    clause, params = scenario_clause(scenario_id)

    exclude_clause, exclude_params = "", ()
    if grant_ids:
        exclude_clause = f"AND grant_id NOT IN ({','.join('?' for _ in grant_ids)})"
        exclude_params = tuple(grant_ids)

    for month in months:
        other_total = db.execute(
            f"""SELECT COALESCE(SUM(percent), 0) AS total FROM allocations
                WHERE student_id = ? AND month = ? {exclude_clause} AND {clause}""",
            (student_id, month, *exclude_params, *params),
        ).fetchone()["total"]
        if other_total + batch_total > 100:
            return f"{student['name']} would be over 100% allocated in {month_label(month)} ({other_total + batch_total}%)."

    for grant_id, percent in rows:
        for month in months:
            db.execute(
                f"DELETE FROM allocations WHERE student_id = ? AND grant_id = ? AND month = ? AND {clause}",
                (student_id, grant_id, month, *params),
            )
            if percent > 0:
                db.execute(
                    "INSERT INTO allocations (scenario_id, student_id, grant_id, month, percent) VALUES (?, ?, ?, ?, ?)",
                    (scenario_id, student_id, grant_id, month, percent),
                )
    return None


def allocations_grouped_by_student(db, scenario_id):
    """{student_id: [{grant_id, grant_name, month, percent}, ...]} for one universe
    (scenario_id=None means live). Used to seed the slider editor's starting point --
    live data on the "New scenario" form, that scenario's own data on its "add a change"
    form, so edits build on whatever's already there rather than always on live.
    """
    clause, params = scenario_clause(scenario_id)
    rows = db.execute(
        f"""SELECT allocations.student_id, allocations.grant_id, grants.name AS grant_name,
                   allocations.month, allocations.percent
            FROM allocations JOIN grants ON grants.id = allocations.grant_id
            WHERE {clause}""",
        params,
    ).fetchall()
    by_student = {}
    for row in rows:
        by_student.setdefault(row["student_id"], []).append(
            {"grant_id": row["grant_id"], "grant_name": row["grant_name"], "month": row["month"], "percent": row["percent"]}
        )
    return by_student


def scenario_changed_student_ids(db, scenario_id):
    """Student ids whose allocation in this scenario differs from live -- added, removed,
    or a changed percent for some grant/month. Used to badge who's actually been touched
    when a scenario spans multiple students.
    """
    live_rows = {
        (r["student_id"], r["grant_id"], r["month"], r["percent"])
        for r in db.execute("SELECT student_id, grant_id, month, percent FROM allocations WHERE scenario_id IS NULL")
    }
    scenario_rows = {
        (r["student_id"], r["grant_id"], r["month"], r["percent"])
        for r in db.execute(
            "SELECT student_id, grant_id, month, percent FROM allocations WHERE scenario_id = ?", (scenario_id,)
        )
    }
    return {row[0] for row in live_rows.symmetric_difference(scenario_rows)}


@app.route("/allocations/set", methods=["POST"])
def set_allocation():
    db = get_db()
    scenario_id = parse_scenario_arg(request.form.get("scenario_id"))
    return_view = request.form.get("return_view")
    return_id = request.form.get("return_id")
    redirect_target = (
        url_for("student_detail", student_id=return_id, scenario=scenario_id)
        if return_view == "student"
        else url_for("grant_detail", grant_id=return_id, scenario=scenario_id)
    )

    error = apply_allocation(
        db,
        scenario_id,
        request.form.get("student_id", ""),
        request.form.get("grant_id", ""),
        request.form.get("month_start", ""),
        request.form.get("month_end", ""),
        request.form.get("percent", ""),
    )
    if error:
        flash(error)
        return redirect(redirect_target)
    db.commit()
    return redirect(redirect_target)


@app.route("/settings")
def settings_page():
    return render_template("settings.html", settings=get_settings(get_db()))


@app.route("/settings/update", methods=["POST"])
def update_settings():
    overspend_bps = parse_rate(request.form.get("overspend_ratio", ""))
    underspend_bps = parse_rate(request.form.get("underspend_ratio", ""))

    if overspend_bps is None or underspend_bps is None:
        flash("Both thresholds must be non-negative percentages.")
    elif underspend_bps >= overspend_bps:
        flash("The underspending threshold must be lower than the overspending threshold.")
    else:
        get_db().execute(
            "UPDATE settings SET overspend_ratio_bps = ?, underspend_ratio_bps = ? WHERE id = 1",
            (overspend_bps, underspend_bps),
        )
        get_db().commit()
    return redirect(url_for("settings_page"))


@app.route("/departments")
def departments_page():
    departments = get_db().execute("SELECT * FROM departments ORDER BY name").fetchall()
    return render_template("departments.html", departments=departments)


@app.route("/departments/add", methods=["POST"])
def add_department():
    name = request.form.get("name", "").strip()
    stipend_cents = parse_money(request.form.get("stipend", "0") or "0")
    tuition_cents = parse_money(request.form.get("tuition", "0") or "0")
    fringe_rate_bps = parse_rate(request.form.get("fringe_rate", "0") or "0")
    tuition_charged_in_summer = 1 if request.form.get("tuition_charged_in_summer") == "on" else 0

    if not name:
        flash("Department name is required.")
    elif stipend_cents is None:
        flash("Stipend must be a non-negative number.")
    elif tuition_cents is None:
        flash("Tuition must be a non-negative number.")
    elif fringe_rate_bps is None:
        flash("Fringe rate must be a non-negative number.")
    else:
        get_db().execute(
            """INSERT INTO departments (name, stipend_cents_per_month, tuition_cents_per_month, fringe_rate_bps,
               tuition_charged_in_summer) VALUES (?, ?, ?, ?, ?)""",
            (name, stipend_cents, tuition_cents, fringe_rate_bps, tuition_charged_in_summer),
        )
        get_db().commit()
    return redirect(url_for("departments_page"))


@app.route("/departments/<int:department_id>/edit", methods=["POST"])
def edit_department(department_id):
    name = request.form.get("name", "").strip()
    stipend_cents = parse_money(request.form.get("stipend", "0") or "0")
    tuition_cents = parse_money(request.form.get("tuition", "0") or "0")
    fringe_rate_bps = parse_rate(request.form.get("fringe_rate", "0") or "0")
    tuition_charged_in_summer = 1 if request.form.get("tuition_charged_in_summer") == "on" else 0

    if not name:
        flash("Department name is required.")
    elif stipend_cents is None:
        flash("Stipend must be a non-negative number.")
    elif tuition_cents is None:
        flash("Tuition must be a non-negative number.")
    elif fringe_rate_bps is None:
        flash("Fringe rate must be a non-negative number.")
    else:
        get_db().execute(
            """UPDATE departments SET name = ?, stipend_cents_per_month = ?, tuition_cents_per_month = ?,
               fringe_rate_bps = ?, tuition_charged_in_summer = ? WHERE id = ?""",
            (name, stipend_cents, tuition_cents, fringe_rate_bps, tuition_charged_in_summer, department_id),
        )
        get_db().commit()
    return redirect(url_for("departments_page"))


@app.route("/departments/<int:department_id>/delete", methods=["POST"])
def delete_department(department_id):
    db = get_db()
    in_use = db.execute("SELECT 1 FROM students WHERE department_id = ? LIMIT 1", (department_id,)).fetchone()
    if in_use:
        flash("Can't delete a department that has students assigned to it. Reassign those students first.")
        return redirect(url_for("departments_page"))
    db.execute("DELETE FROM departments WHERE id = ?", (department_id,))
    db.commit()
    return redirect(url_for("departments_page"))


@app.route("/scenarios")
def scenarios_page():
    db = get_db()
    scenarios = db.execute("SELECT * FROM scenarios ORDER BY created_at DESC").fetchall()
    students = db.execute("SELECT * FROM students ORDER BY name").fetchall()
    grants = db.execute("SELECT * FROM grants ORDER BY name").fetchall()
    grants_json = [{"id": g["id"], "name": g["name"]} for g in grants]

    return render_template(
        "scenarios.html",
        scenarios=scenarios,
        students=students,
        grants=grants,
        grants_json=grants_json,
        today=date.today().isoformat(),
        baseline_allocations_by_student=allocations_grouped_by_student(db, None),
    )


@app.route("/scenarios/plan-new-student")
def plan_new_student():
    """Read-only capacity report: given a time window and a set of hypothetical new
    hires (by department + role), checks whether current grants have enough spare
    balance to cover them and shows 2-3 candidate ways to split their effort. Doesn't
    write anything -- no student, allocation, or scenario is created. GET (not POST)
    since it's a pure query with no side effects.
    """
    db = get_db()
    departments = db.execute("SELECT * FROM departments ORDER BY name").fetchall()

    window_start_raw = request.args.get("window_start", "")
    window_end_raw = request.args.get("window_end", "")
    department_ids = request.args.getlist("department_id[]")
    roles = request.args.getlist("role[]")
    counts = request.args.getlist("count[]")

    plans = None
    excluded_grants = []
    hire_rows_prefill = []
    submitted = bool(window_start_raw or window_end_raw or department_ids)

    if submitted:
        window_start = parse_month(window_start_raw)
        window_end = parse_month(window_end_raw)
        if window_start is None or window_end is None:
            flash("Start and end month must be valid.")
        elif window_end < window_start:
            flash("End month can't be before the start month.")
        else:
            months = month_range(window_start, window_end)
            hires = []
            for dept_id_raw, role_raw, count_raw in zip(department_ids, roles, counts):
                dept_row = next((d for d in departments if d["id"] == parse_department_id(dept_id_raw)), None)
                role = parse_role(role_raw)
                try:
                    count = int(count_raw)
                except (TypeError, ValueError):
                    count = 0
                if dept_row is None:
                    flash("Skipped a row: that department no longer exists.")
                    continue
                if role is None or count <= 0:
                    continue
                dept = dict(dept_row)
                hire_rows_prefill.append({"department_id": dept["id"], "department_name": dept["name"], "role": role, "count": count})
                hires.extend({"department": dept, "role": role} for _ in range(count))

            if not hires:
                flash("Add at least one department/role/count row.")
            else:
                eligible_grants, excluded_grants = eligible_grants_for_window(db, months)
                plans = [
                    build_greedy_hire_plan(
                        "Concentrate",
                        "Fills each grant to capacity, starting with the grant with the most spare capacity, before moving to the next.",
                        hires,
                        sorted(eligible_grants, key=lambda g: -g["spare_cents"]),
                        months,
                    ),
                    build_greedy_hire_plan(
                        "Use it or lose it",
                        "Same as Concentrate, but prioritizes grants that expire soonest so their balance gets used before it's lost.",
                        hires,
                        sorted(eligible_grants, key=lambda g: g["end_date"]),
                        months,
                    ),
                    build_spread_hire_plan(hires, eligible_grants, months),
                ]

    return render_template(
        "plan_new_student.html",
        departments=departments,
        departments_json=[{"id": d["id"], "name": d["name"]} for d in departments],
        window_start=window_start_raw,
        window_end=window_end_raw,
        hire_rows_prefill=hire_rows_prefill,
        plans=plans,
        excluded_grants=excluded_grants,
        today=date.today().isoformat(),
    )


@app.route("/scenarios/add", methods=["POST"])
def add_scenario():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Scenario name is required.")
        return redirect(url_for("scenarios_page"))

    db = get_db()
    cursor = db.execute(
        "INSERT INTO scenarios (name, created_at) VALUES (?, ?)",
        (name, datetime.now().isoformat(timespec="seconds")),
    )
    scenario_id = cursor.lastrowid
    db.execute(
        """INSERT INTO allocations (scenario_id, student_id, grant_id, month, percent)
           SELECT ?, student_id, grant_id, month, percent FROM allocations WHERE scenario_id IS NULL""",
        (scenario_id,),
    )

    student_id = request.form.get("student_id", "")
    grant_ids = request.form.getlist("grant_id[]")
    percents = request.form.getlist("percent[]")
    rows = [(g, p) for g, p in zip(grant_ids, percents) if g]
    if student_id and rows:
        error = apply_allocation_batch(
            db,
            scenario_id,
            student_id,
            rows,
            request.form.get("month_start", ""),
            request.form.get("month_end", ""),
        )
        if error:
            db.rollback()
            flash(error)
            return redirect(url_for("scenarios_page"))

    db.commit()
    return redirect(url_for("scenario_detail", scenario_id=scenario_id))


@app.route("/scenarios/<int:scenario_id>")
def scenario_detail(scenario_id):
    db = get_db()
    scenario = db.execute("SELECT * FROM scenarios WHERE id = ?", (scenario_id,)).fetchone()
    if scenario is None:
        return redirect(url_for("scenarios_page"))

    live_student_ids = {r["student_id"] for r in db.execute("SELECT DISTINCT student_id FROM allocations WHERE scenario_id IS NULL")}
    scenario_student_ids = {
        r["student_id"] for r in db.execute("SELECT DISTINCT student_id FROM allocations WHERE scenario_id = ?", (scenario_id,))
    }
    changed_student_ids = scenario_changed_student_ids(db, scenario_id)

    student_rows = []
    for student_id in live_student_ids | scenario_student_ids:
        student = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
        if student is None:
            continue
        live_grid = student_allocation_grid(student_id, None)
        scenario_grid = student_allocation_grid(student_id, scenario_id)
        student_rows.append(
            {
                "id": student_id,
                "name": student["name"],
                "changed": student_id in changed_student_ids,
                "live_cost_cents": live_grid["breakdown"]["total"] if live_grid else 0,
                "scenario_cost_cents": scenario_grid["breakdown"]["total"] if scenario_grid else 0,
                "under_allocated_count": len(under_allocated_months(scenario_grid)),
            }
        )
    student_rows.sort(key=lambda r: r["name"])

    live_grant_ids = {r["grant_id"] for r in db.execute("SELECT DISTINCT grant_id FROM allocations WHERE scenario_id IS NULL")}
    scenario_grant_ids = {
        r["grant_id"] for r in db.execute("SELECT DISTINCT grant_id FROM allocations WHERE scenario_id = ?", (scenario_id,))
    }
    grant_rows = []
    for grant_id in live_grant_ids | scenario_grant_ids:
        grant = db.execute("SELECT * FROM grants WHERE id = ?", (grant_id,)).fetchone()
        if grant is None:
            continue
        live_grid = grant_allocation_grid(grant_id, None)
        scenario_grid = grant_allocation_grid(grant_id, scenario_id)
        grant_rows.append(
            {
                "id": grant_id,
                "name": grant["name"],
                "live_cost_cents": live_grid["total_cost"] if live_grid else 0,
                "scenario_cost_cents": scenario_grid["total_cost"] if scenario_grid else 0,
            }
        )
    grant_rows.sort(key=lambda r: r["name"])

    students = db.execute("SELECT * FROM students ORDER BY name").fetchall()
    grants = db.execute("SELECT * FROM grants ORDER BY name").fetchall()
    grants_json = [{"id": g["id"], "name": g["name"]} for g in grants]

    return render_template(
        "scenario_detail.html",
        scenario=scenario,
        student_rows=student_rows,
        grant_rows=grant_rows,
        total_live_cents=sum(r["live_cost_cents"] for r in grant_rows),
        total_scenario_cents=sum(r["scenario_cost_cents"] for r in grant_rows),
        students=students,
        grants=grants,
        grants_json=grants_json,
        today=date.today().isoformat(),
        baseline_allocations_by_student=allocations_grouped_by_student(db, scenario_id),
    )


@app.route("/scenarios/<int:scenario_id>/allocations/add", methods=["POST"])
def add_scenario_allocation(scenario_id):
    db = get_db()
    scenario = db.execute("SELECT id FROM scenarios WHERE id = ?", (scenario_id,)).fetchone()
    if scenario is None:
        return redirect(url_for("scenarios_page"))

    student_id = request.form.get("student_id", "")
    grant_ids = request.form.getlist("grant_id[]")
    percents = request.form.getlist("percent[]")
    rows = [(g, p) for g, p in zip(grant_ids, percents) if g]
    if not student_id or not rows:
        flash("Select a student and at least one grant.")
        return redirect(url_for("scenario_detail", scenario_id=scenario_id))

    error = apply_allocation_batch(
        db, scenario_id, student_id, rows, request.form.get("month_start", ""), request.form.get("month_end", "")
    )
    if error:
        flash(error)
        return redirect(url_for("scenario_detail", scenario_id=scenario_id))
    db.commit()
    return redirect(url_for("scenario_detail", scenario_id=scenario_id))


@app.route("/scenarios/<int:scenario_id>/apply", methods=["POST"])
def apply_scenario(scenario_id):
    db = get_db()
    db.execute("DELETE FROM allocations WHERE scenario_id IS NULL")
    db.execute(
        """INSERT INTO allocations (scenario_id, student_id, grant_id, month, percent)
           SELECT NULL, student_id, grant_id, month, percent FROM allocations WHERE scenario_id = ?""",
        (scenario_id,),
    )
    db.commit()
    flash("Scenario applied to live data.")
    return redirect(url_for("scenarios_page"))


@app.route("/scenarios/<int:scenario_id>/delete", methods=["POST"])
def delete_scenario(scenario_id):
    db = get_db()
    db.execute("DELETE FROM allocations WHERE scenario_id = ?", (scenario_id,))
    db.execute("DELETE FROM scenarios WHERE id = ?", (scenario_id,))
    db.commit()
    return redirect(url_for("scenarios_page"))


@app.route("/faculty")
def faculty_page():
    faculty = [{"slug": slug, "name": faculty_display_name(slug)} for slug in list_faculty()]
    return render_template("faculty.html", faculty=faculty, current=session.get("faculty_db"))


@app.route("/faculty/select", methods=["POST"])
def select_faculty():
    slug = request.form.get("slug", "")
    if slug not in list_faculty():
        flash("Select a valid faculty database.")
        return redirect(url_for("faculty_page"))
    session["faculty_db"] = slug
    return redirect(url_for("index"))


@app.route("/faculty/add", methods=["POST"])
def add_faculty():
    name = request.form.get("name", "").strip()
    slug = slugify(name)
    if not name or not slug:
        flash("Enter a faculty name.")
    elif slug in list_faculty():
        flash("A database for that name already exists. Select it from the list instead.")
    else:
        init_db(faculty_db_path(slug))
        session["faculty_db"] = slug

        workbook_file = request.files.get("workbook")
        if workbook_file and workbook_file.filename:
            import_grants_from_workbook(workbook_file)

        return redirect(url_for("index"))
    return redirect(url_for("faculty_page"))


def import_grants_from_workbook(workbook_file):
    """Best-effort import of active grants from an uploaded .xlsx report into the
    just-created, currently-active faculty db (see add_faculty()). Reuses the same
    header-detection/parsing logic as import_from_excel.py's CLI, with its default
    options (today as as-of-date, 3-year start-date estimate, excludes expired
    awards). Failures here are non-fatal -- the faculty db was already created and
    stays around either way, flash() just tells the user what happened.
    """
    from import_from_excel import REPORT_REQUIRED_COLUMNS, build_grant_records, existing_ptas, find_header

    try:
        wb = openpyxl.load_workbook(BytesIO(workbook_file.read()), data_only=True, read_only=True)
    except Exception:
        flash("Couldn't read that file as an Excel workbook. The faculty database was still created.")
        return

    sheet_name, header_row, col = find_header(wb, REPORT_REQUIRED_COLUMNS)
    if sheet_name is None:
        flash(
            "Couldn't find a matching report sheet in that workbook. The faculty database was still "
            "created -- add grants by hand, or try importing again later from the Faculty page."
        )
        return

    today = date.today()
    records = build_grant_records(wb, sheet_name, col, today, include_expired=False, estimate_years=3)
    db = get_db()
    already = existing_ptas(db)
    inserted = 0
    for r in records:
        if r["pta"] in already:
            continue
        cursor = db.execute(
            """INSERT INTO grants (name, sponsor, total_amount_cents, start_date, end_date, overhead_rate_bps, notes)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (r["name"], "", r["total_amount_cents"], r["start_date"], r["end_date"], 0, r["notes"]),
        )
        if r["spent_cents"]:
            db.execute(
                "INSERT INTO transactions (grant_id, date, amount_cents, description) VALUES (?, ?, ?, ?)",
                (
                    cursor.lastrowid,
                    today.isoformat(),
                    r["spent_cents"],
                    f"Expenditures + encumbrances as of {today.isoformat()} (imported from {workbook_file.filename})",
                ),
            )
        inserted += 1
    db.commit()
    flash(f"Imported {inserted} grant(s) from {workbook_file.filename}.")


if __name__ == "__main__":
    import os

    # debug=True is safe only because app.run() binds to 127.0.0.1 by default.
    # If you ever pass host="0.0.0.0" to expose this beyond localhost, turn
    # debug off first -- the interactive debugger allows arbitrary code
    # execution to anyone who can reach this port.
    app.run(
        debug=True,
        port=int(os.environ.get("PORT", 5000)),
        use_reloader=os.environ.get("DISABLE_RELOADER") != "1",
    )
